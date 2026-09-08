from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "deliverables" / "南联青训正式发布前最终验收流程表.docx"


PHASES = [
    ("第一阶段  封版与测试准备", [
        ("1.1", "确定测试负责人", "填写总负责人、管理员测试人、教练测试人和家长测试人。"),
        ("1.2", "确认当前体验版", "微信体验版为 1.0.1，能够正常扫码打开。"),
        ("1.3", "准备真实测试账号", "至少准备 1 个 ADMIN、1 个已绑定 COACH、2 个 PARENT 微信账号。"),
        ("1.4", "准备多孩子家庭", "其中一个家长账号至少关联 2 个测试孩子。"),
        ("1.5", "准备班级", "建立 1 个普通班和 1 个精英队，并明确年龄段、教练、时间、场地。"),
        ("1.6", "准备测试课程", "发布至少 2 节未来课程，其中 1 节用于教练真实点名。"),
        ("1.7", "准备成长数据", "准备 1 条周训练计划、1 条阶段评价和 1 条成长记录。"),
        ("1.8", "准备问题记录方式", "所有问题记录角色、页面、操作步骤、期望结果、实际结果和截图。"),
    ]),
    ("第二阶段  管理员真实验收", [
        ("2.1", "管理员登录", "已有 ADMIN 自动进入管理员端，陌生微信不能获得管理员权限。"),
        ("2.2", "审核孩子资料", "运营后台能够找到孩子资料审核入口，能查看并处理待审核资料。"),
        ("2.3", "核对学员照片", "审核后照片、姓名、出生日期、学校、年级正确保存和显示。"),
        ("2.4", "核对身份证权限", "完整身份证只在管理员授权页面显示，普通接口不返回完整号码。"),
        ("2.5", "核对家长归属", "一个家长可有多个孩子，一个孩子只能归属一个主家长。"),
        ("2.6", "测试重复身份证", "用同一身份证再次创建孩子时必须识别重复并阻止创建。"),
        ("2.7", "创建和编辑班级", "班级名称、类型、年龄段、训练时间、场地和教练均可正确维护。"),
        ("2.8", "添加正式成员", "添加后生成 ACTIVE classMember，家长立即能看到我的班级。"),
        ("2.9", "创建和发布课程", "发布后班级正式成员自动在课表中看到课程，无需逐节报名。"),
        ("2.10", "调整单节课程", "临时更换时间、场地或教练后，三端显示一致。"),
        ("2.11", "取消课程", "取消课程不扣课；补偿或补课逻辑按管理员选择执行且不可重复。"),
        ("2.12", "生成教练邀请", "管理员可为指定教练档案生成绑定邀请，邀请内容与教练一致。"),
        ("2.13", "审核精英推荐", "教练推荐只进入待审核，最终通过、驳回和晋升由管理员完成。"),
    ]),
    ("第三阶段  教练真实验收", [
        ("3.1", "绑定真实教练微信", "教练通过二维码或备用邀请码完成绑定，自动关联正确 coachProfile。"),
        ("3.2", "教练自动识别", "再次进入时自动识别 COACH，不要求重复选择身份。"),
        ("3.3", "核对教练首页", "只显示今日教学、我的班级、学员成长、推荐精英队、我的工作。"),
        ("3.4", "核对今日教学", "能够看到本人今天、本周及全部已安排课程。"),
        ("3.5", "打开点名页面", "今日教学 → 开始点名，显示本人负责班级的正式学员照片和姓名。"),
        ("3.6", "测试四种考勤状态", "分别保存到课、请假、伤病、缺勤，刷新后记录仍存在。"),
        ("3.7", "测试全部到课", "一键全部到课可用，但不会覆盖已经锁定的批准请假。"),
        ("3.8", "测试未点完保护", "存在未点名学员时不能提交，并明确提示剩余人数。"),
        ("3.9", "核对消课规则", "到课、缺勤扣 1 节；请假、伤病不扣；重复保存不重复扣课。"),
        ("3.10", "填写课后记录", "能够从课程进入课后记录，并保存允许范围内的训练评价。"),
        ("3.11", "新增阶段评价", "学员成长 → 选择本人学员 → 新增评价，保存后家长可见公开内容。"),
        ("3.12", "提交精英推荐", "选择本人负责班级 ACTIVE 学员，填写理由并提交成功。"),
        ("3.13", "跨班越权测试", "修改 classId、studentId 或 coachId 不能读取或修改其他教练数据。"),
        ("3.14", "敏感功能隔离", "教练不能访问订单、支付、CRM、俱乐部收入或管理员系统。"),
    ]),
    ("第四阶段  家长真实验收", [
        ("4.1", "陌生微信首次进入", "进入欢迎登录注册页，不默认成为 PARENT，不显示待绑定家长。"),
        ("4.2", "注册家长", "填写家长姓名和手机号后成功创建 PARENT，授权取消有明确提示。"),
        ("4.3", "无孩子空状态", "显示欢迎加入南联和添加孩子入口，不显示错误身份提示。"),
        ("4.4", "添加第一个孩子", "照片、姓名、性别、出生日期、身份证、学校、年级可提交审核。"),
        ("4.5", "添加第二个孩子", "同一家长可以继续添加孩子，两个孩子资料互不覆盖。"),
        ("4.6", "资料审核后可见", "管理员通过后，孩子出现在我的孩子且归属正确。"),
        ("4.7", "完成班级报名", "报名成功后立即显示我的班级和正式成员状态。"),
        ("4.8", "无未来课程提示", "已经入班但未排课时，明确提示已入班且训练安排暂未发布。"),
        ("4.9", "查看我的课程", "显示下一次训练、当前班级、本周及近期课程。"),
        ("4.10", "多孩子课程切换", "全部孩子和单个孩子均可查看，课程、班级与姓名不串数据。"),
        ("4.11", "提交请假", "请假页直接显示最近可请假课程，可完成提交和查看状态。"),
        ("4.12", "核对请假结果", "批准后课程和班级仍存在，本节显示已请假且不扣课。"),
        ("4.13", "查看成长", "只能看到本人孩子的公开评价、成长记录和阶段测评。"),
        ("4.14", "查看课时与购课", "余额、套餐和订单显示正确；支付关闭时提示微信支付暂未开通。"),
        ("4.15", "家长越权测试", "直接替换 studentId 不能读取其他家长孩子的资料、照片或成长信息。"),
    ]),
    ("第五阶段  业务红线与安全", [
        ("5.1", "固定班级成员", "classMembers 是正式班级成员唯一权威来源。"),
        ("5.2", "候补机制关闭", "系统不存在候补、自动补位或请假释放名额。"),
        ("5.3", "体验课隔离", "体验学员不是正式 classMember，体验课不扣 lessonLedger。"),
        ("5.4", "请假边界", "请假只影响单节 attendance，不退出班级、不自动延期套餐。"),
        ("5.5", "精英队规则", "教练推荐加管理员审核，不按照测评分数自动晋升。"),
        ("5.6", "成长联赛规则", "比赛不扣训练课时，外部球员不创建正式 student。"),
        ("5.7", "支付集合权限", "orders、payments、权益和流水集合禁止小程序客户端直接写入。"),
        ("5.8", "身份权限云端校验", "修改前端 role 不能获得 COACH 或 ADMIN 权限。"),
        ("5.9", "账号停用", "停用账号不能进入业务首页，历史数据不被删除。"),
        ("5.10", "敏感数据白名单", "PARENT 和 COACH 普通接口不返回完整身份证、openid 或支付信息。"),
        ("5.11", "测试角色入口关闭", "清空 TEST_ROLE_SWITCH_OPENIDS，移除硬编码测试哈希并隐藏测试角色入口。"),
        ("5.12", "支付开关确认", "未完成 1 元真实支付验收前，PAYMENT_PRODUCTION_ENABLED 明确为 false。"),
    ]),
    ("第六阶段  隐私与平台配置", [
        ("6.1", "用户隐私保护指引", "微信公众平台已配置并通过，内容与程序实际处理的信息一致。"),
        ("6.2", "手机号说明", "说明收集家长手机号的用途、使用范围和必要性。"),
        ("6.3", "未成年人资料说明", "说明姓名、出生日期、学校、年级等资料用于培训管理。"),
        ("6.4", "照片说明", "说明孩子照片仅用于内部点名和身份识别，不默认公开宣传。"),
        ("6.5", "身份证说明", "说明身份核验用途、访问权限、保存和删除方式。"),
        ("6.6", "定位说明", "说明定位仅用于课程签到；拒绝授权时可以返回或使用教练点名。"),
        ("6.7", "授权失败处理", "相机、相册、手机号和定位拒绝后均有明确提示，不出现死页面。"),
        ("6.8", "客服与主体信息", "小程序名称、公司主体、客服电话和联系渠道真实有效。"),
        ("6.9", "服务类目", "公众平台服务类目与青少年足球培训实际业务一致。"),
    ]),
    ("第七阶段  技术封版与云端", [
        ("7.1", "备份 CloudBase 数据", "发布前导出或备份核心集合，并记录备份时间和负责人。"),
        ("7.2", "核对正式环境", "小程序使用 cloud1-d2g4gi77g48dcee01，不连接本地或错误环境。"),
        ("7.3", "部署最终 clubApi", "关闭测试角色能力后，将最终云函数部署到 cloud1 并记录时间。"),
        ("7.4", "检查云函数日志", "登录、课程、点名、请假和成长接口没有持续异常。"),
        ("7.5", "运行全部自动化测试", "全部 tests/*.js 通过，不为了得到 PASS 修改测试规则。"),
        ("7.6", "检查页面完整性", "Registered pages 71，Missing page files 0。"),
        ("7.7", "检查主包大小", "主包低于微信 2MB 限制。"),
        ("7.8", "检查 Git 差异", "git status、git diff 和 git diff --check 均无不明修改或删除。"),
        ("7.9", "合并 PR 到 main", "PR #12 审查通过并合并，正式发布代码以主仓库 main 为准。"),
        ("7.10", "创建正式版本标识", "为最终 main 创建 v1.0.0 或公司确认的正式版本标签。"),
        ("7.11", "上传最终候选版", "从干净 main 上传新候选版本，不使用仍在开发的分支直接提审。"),
        ("7.12", "复测最终候选版", "上传后再次完成登录、家长课程、教练点名和管理员学员管理冒烟测试。"),
    ]),
    ("第八阶段  提审发布与观察", [
        ("8.1", "填写审核说明", "清楚说明家长、教练和管理员的业务用途及账号授权方式。"),
        ("8.2", "提交微信审核", "公众平台 → 管理 → 版本管理 → 开发版本 → 提交审核。"),
        ("8.3", "处理审核反馈", "记录审核退回原因，只做对应修复，重新回归后再提交。"),
        ("8.4", "选择发布时间", "安排在管理员和技术人员在线、非训练高峰的工作日上午。"),
        ("8.5", "执行正式发布", "审核通过后由管理员点击发布，并记录版本号与发布时间。"),
        ("8.6", "发布后家长冒烟", "家长登录、孩子、班级、课程、请假和成长页面正常。"),
        ("8.7", "发布后教练冒烟", "教练登录、今日教学、开始点名和保存考勤正常。"),
        ("8.8", "发布后管理员冒烟", "管理员登录、孩子审核、班级课程和学员管理正常。"),
        ("8.9", "观察 24 小时", "持续检查云函数错误、权限异常、报名、点名、课时和支付状态。"),
        ("8.10", "形成首日结论", "填写是否稳定运营、遗留问题、负责人和预计处理时间。"),
    ]),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=55, start=95, bottom=55, end=95):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, bold=False, color="000000", size=10, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial Unicode MS"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def add_check_table(doc, items):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(0.52), Inches(1.50), Inches(4.36), Inches(0.96)]
    headers = ["序号", "任务", "验收标准", "结果"]
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, cell in enumerate(header.cells):
        cell.width = widths[idx]
        set_cell_margins(cell, 75, 95, 75, 95)
        set_cell_shading(cell, "173B77")
        set_cell_text(cell, headers[idx], bold=True, color="FFFFFF", size=9.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row_index, (number, task, standard) in enumerate(items):
        cells = table.add_row().cells
        for idx, cell in enumerate(cells):
            cell.width = widths[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, "F4F7FB")
        set_cell_text(cells[0], number, size=8.8, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(cells[1], task, bold=True, size=9.0)
        set_cell_text(cells[2], standard, size=9.0)
        set_cell_text(cells[3], "□通过\n□未通过", size=8.0, align=WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.58)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.58)
section.right_margin = Inches(0.58)
section.header_distance = Inches(0.25)
section.footer_distance = Inches(0.25)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial Unicode MS"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.15

title_style = styles["Title"]
title_style.font.name = "Arial Unicode MS"
title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
title_style.font.size = Pt(24)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)

for style_name, size in (("Heading 1", 16), ("Heading 2", 13)):
    style = styles[style_name]
    style.font.name = "Arial Unicode MS"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(7)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = header.add_run("永嘉南联体育培训有限公司  正式发布验收")
run.font.name = "Arial Unicode MS"
run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
run.font.size = Pt(8.5)
run.font.color.rgb = RGBColor(90, 90, 90)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("第 ")
run.font.size = Pt(8.5)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)
run = footer.add_run(" 页")
run.font.size = Pt(8.5)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("南联青训正式发布前最终验收流程表")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("适用于微信小程序体验版验收  提审  正式发布及首日观察")
r.font.name = "Arial Unicode MS"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(75, 75, 75)

meta = doc.add_table(rows=4, cols=4)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
meta.autofit = False
meta_widths = [Inches(1.25), Inches(2.4), Inches(1.25), Inches(2.4)]
meta_rows = [
    ("公司", "永嘉南联体育培训有限公司", "小程序", "南联青训云端管理版"),
    ("当前体验版", "1.0.1", "当前提交", "82850c4"),
    ("总负责人", "________________", "计划发布日期", "______年____月____日"),
    ("最终结论", "□ 允许提审  □ 暂停", "签字", "________________"),
]
for row_index, values in enumerate(meta_rows):
    for col_index, value in enumerate(values):
        cell = meta.cell(row_index, col_index)
        cell.width = meta_widths[col_index]
        set_cell_margins(cell, 125, 120, 125, 120)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if col_index in (0, 2):
            set_cell_shading(cell, "E8EEF7")
            set_cell_text(cell, value, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            set_cell_text(cell, value, size=10)

doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("使用规则  ")
r.bold = True
r.font.size = Pt(11)
p.add_run("按顺序逐项测试。任何 P0 问题、权限越权、数据错账、身份识别错误或点名消课错误出现时，立即停止提审；修复后必须重新执行关联测试和全部自动化回归。")

doc.add_heading("当前已知基线", level=1)
baseline = [
    "体验版 1.0.1 已上传，尚未提交正式审核。",
    "本地 48 个自动化测试文件已通过；正式候选版仍需再次完整回归。",
    "Registered pages 为 71，Missing page files 为 0，主包约 702.9KB。",
    "微信支付生产开关应继续保持关闭，直到 1 元真实支付验收完成。",
    "正式发布前必须关闭测试角色切换白名单和硬编码测试账号能力。",
]
for item in baseline:
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(item)

for phase_index, (phase_title, items) in enumerate(PHASES):
    heading = doc.add_heading(phase_title, level=1)
    heading.paragraph_format.page_break_before = True
    p = doc.add_paragraph()
    p.add_run(f"本阶段共 {len(items)} 项。完成一项后立即勾选结果；未通过项目必须在问题记录表登记。").italic = True
    add_check_table(doc, items)
    p = doc.add_paragraph()
    p.add_run("阶段结论  ").bold = True
    p.add_run("□ 全部通过    □ 有条件通过    □ 暂停进入下一阶段")
    p = doc.add_paragraph()
    p.add_run("负责人签字  ").bold = True
    p.add_run("________________________    日期  ______年____月____日")

heading = doc.add_heading("问题记录表", level=1)
heading.paragraph_format.page_break_before = True
p = doc.add_paragraph("发现问题时先记录，不在未确认影响范围前进行大范围修改。P0 问题必须关闭后才能提审。")
issue_table = doc.add_table(rows=1, cols=6)
issue_table.alignment = WD_TABLE_ALIGNMENT.CENTER
issue_table.autofit = False
issue_widths = [Inches(0.55), Inches(0.72), Inches(1.0), Inches(2.35), Inches(1.45), Inches(1.0)]
issue_headers = ["编号", "级别", "角色", "问题与复现步骤", "负责人及期限", "复测结果"]
for idx, cell in enumerate(issue_table.rows[0].cells):
    cell.width = issue_widths[idx]
    set_cell_margins(cell)
    set_cell_shading(cell, "173B77")
    set_cell_text(cell, issue_headers[idx], bold=True, color="FFFFFF", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
set_repeat_table_header(issue_table.rows[0])
for row_no in range(1, 11):
    cells = issue_table.add_row().cells
    for idx, cell in enumerate(cells):
        cell.width = issue_widths[idx]
        set_cell_margins(cell, 170, 100, 170, 100)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if row_no % 2 == 0:
            set_cell_shading(cell, "F4F7FB")
    values = [str(row_no), "□P0\n□P1\n□P2", "", "", "", "□通过\n□未通过"]
    for idx, value in enumerate(values):
        set_cell_text(cells[idx], value, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER if idx != 3 else WD_ALIGN_PARAGRAPH.LEFT)

heading = doc.add_heading("最终发布批准", level=1)
heading.paragraph_format.page_break_before = True
final_items = [
    "□ ADMIN 真实账号验收通过",
    "□ COACH 真实账号与点名消课验收通过",
    "□ PARENT 注册、添加孩子、报名、课程和请假验收通过",
    "□ 测试角色切换能力已关闭并重新部署",
    "□ 隐私保护指引、服务类目和客服联系信息已确认",
    "□ 支付保持已批准的安全状态",
    "□ CloudBase 数据已备份，云函数日志无持续异常",
    "□ 全量自动化测试、页面完整性和主包大小检查通过",
    "□ PR 已合并 main，最终候选版从干净 main 上传",
    "□ 当前不存在 P0 问题",
]
for item in final_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(item)
    r.font.size = Pt(11)

approval = doc.add_table(rows=5, cols=2)
approval.alignment = WD_TABLE_ALIGNMENT.CENTER
approval.autofit = False
approval_data = [
    ("最终版本号", "____________________________"),
    ("微信审核提交时间", "______年____月____日  ____时____分"),
    ("微信审核结果", "□ 通过    □ 退回    □ 审核中"),
    ("正式发布时间", "______年____月____日  ____时____分"),
    ("负责人批准签字", "____________________________"),
]
for row_index, (label, value) in enumerate(approval_data):
    for col_index, text in enumerate((label, value)):
        cell = approval.cell(row_index, col_index)
        cell.width = Inches(2.0 if col_index == 0 else 5.0)
        set_cell_margins(cell, 150, 130, 150, 130)
        if col_index == 0:
            set_cell_shading(cell, "E8EEF7")
            set_cell_text(cell, text, bold=True, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            set_cell_text(cell, text, size=10.5)

doc.core_properties.title = "南联青训正式发布前最终验收流程表"
doc.core_properties.subject = "微信小程序体验版验收 提审 发布与首日观察"
doc.core_properties.author = "永嘉南联体育培训有限公司"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
